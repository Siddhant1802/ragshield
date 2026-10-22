from pathlib import Path

from langchain_core.documents import Document

SUPPORTED = {".txt", ".md", ".pdf", ".docx"}


def load_file(path: Path, trust: str = "untrusted") -> list[Document]:
    """Read one file into Document objects. Every document is labelled with its trust level.
    Default is 'untrusted': a source must be deliberately marked as trusted."""
    suffix = path.suffix.lower()
    meta = {"source": path.name, "trust": trust}

    if suffix in {".txt", ".md"}:
        text = path.read_text(encoding="utf-8", errors="ignore")
        return [Document(page_content=text, metadata=dict(meta))]

    if suffix == ".pdf":
        from pypdf import PdfReader
        reader = PdfReader(path)
        # One Document per page, so answers can cite the page number
        return [
            Document(page_content=page.extract_text() or "", metadata={**meta, "page": i + 1})
            for i, page in enumerate(reader.pages)
        ]

    if suffix == ".docx":
        import docx
        d = docx.Document(path)
        text = "\n".join(p.text for p in d.paragraphs)
        return [Document(page_content=text, metadata=dict(meta))]

    return []  # unsupported file type: skip it


def load_documents(folder: str, trust: str = "untrusted") -> list[Document]:
    """Load every supported file in a folder (including subfolders), all with the same trust level."""
    docs = []
    for path in sorted(Path(folder).rglob("*")):
        if path.suffix.lower() in SUPPORTED:
            docs.extend(load_file(path, trust=trust))
    # Drop empty documents, e.g. scanned PDF pages with no extractable text
    return [d for d in docs if d.page_content.strip()]